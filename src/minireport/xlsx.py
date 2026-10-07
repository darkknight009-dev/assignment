"""Build the XLSX results workbook from the raw CSV results."""

from __future__ import annotations

from pathlib import Path

from .models import DRAFT_BANNER, Evidence, Protocol


def build_workbook(evidence: Evidence, protocol: Protocol, out_path: Path) -> None:
    """Create results_table.xlsx: data, limits, deterministic Pass/Fail, summary."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    red_font = Font(color="9C0006", bold=True)
    gray = Font(color="666666", italic=True, size=9)
    bold = Font(bold=True)

    # ---------------------------------------------------------------- data ---
    ws = wb.active
    ws.title = "Results"
    ws["A1"] = DRAFT_BANNER
    ws["A1"].font = Font(bold=True, color="9C0006", size=12)
    ws["A2"] = (f"Protocol {protocol.protocol_id}: {protocol.component} -- evaluated against "
                f"[{protocol.lower_limit}, {protocol.upper_limit}] {protocol.units}")
    ws["A2"].font = gray

    headers = ["Sample ID", "Measurement", "Unit", "Observation", "Acceptance limits",
               "Lower limit", "Upper limit", "Pass/Fail", "Status detail", "CSV row"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = bold

    for i, ev in enumerate(evidence.evaluations, start=5):
        ws.cell(row=i, column=1, value=ev["sample_id"])
        m = ws.cell(row=i, column=2, value=ev["measurement"])
        m.number_format = "0.000"
        ws.cell(row=i, column=3, value=ev["unit"])
        ws.cell(row=i, column=4, value=ev["observation"])
        ws.cell(row=i, column=5, value=_limits_text(protocol))
        ws.cell(row=i, column=6, value=protocol.lower_limit)
        ws.cell(row=i, column=7, value=protocol.upper_limit)
        verdict = "NO DATA" if ev["verdict"] == "NO DATA" else ev["verdict"]
        vcell = ws.cell(row=i, column=8, value=verdict)
        detail = ev["detail"] or ""
        ws.cell(row=i, column=9, value=detail)
        ws.cell(row=i, column=10, value=ev["row_number"])
        if verdict == "Fail":
            for col in range(1, 11):
                ws.cell(row=i, column=col).fill = red_fill
            vcell.font = red_font
        elif verdict == "NO DATA":
            vcell.font = Font(color="996600", bold=True)

    width = [12, 12, 8, 28, 16, 10, 10, 10, 40, 8]
    for col, w in enumerate(width, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A5"

    # ------------------------------------------------------------- summary ---
    st = wb.create_sheet("Summary")
    st["A1"] = DRAFT_BANNER
    st["A1"].font = Font(bold=True, color="9C0006", size=12)
    st["A3"] = "Metric"
    st["B3"] = "Value"
    st["A3"].font = bold
    st["B3"].font = bold
    rows = [
        ("Protocol ID", protocol.protocol_id),
        ("Component", protocol.component),
        ("Total samples", evidence.total),
        ("Passed", evidence.passed),
        ("Failed", evidence.failed),
        ("No data", sum(1 for e in evidence.evaluations if e["verdict"] == "NO DATA")),
        ("Acceptance range", _limits_text(protocol)),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        st.cell(row=i, column=1, value=k)
        st.cell(row=i, column=2, value=v)

    wb.save(str(out_path))


def _limits_text(protocol: Protocol) -> str:
    low = "open" if protocol.lower_limit is None else f"{protocol.lower_limit:g}"
    high = "open" if protocol.upper_limit is None else f"{protocol.upper_limit:g}"
    return f"{low} to {high} {protocol.units}"
