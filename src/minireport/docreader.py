"""Extract plain text from the supported reference/attachment formats."""

from __future__ import annotations

from pathlib import Path


def read_any(path: str | Path) -> list[tuple[int | None, str]]:
    """Return a list of (page_number_or_None, text) blocks.

    * ``.txt`` / ``.md`` -> one block, page None
    * ``.docx``          -> one block per document, page None
    * ``.pdf``           -> one block per page, page number 1-based
    """
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in {".txt", ".md"}:
        return [(None, p.read_text(encoding="utf-8", errors="replace"))]
    if suffix == ".docx":
        return [(None, _read_docx(p))]
    if suffix == ".pdf":
        return _read_pdf_pages(p)
    raise ValueError(f"Unsupported document type: {suffix} ({p.name})")


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [para.text for para in doc.paragraphs if para.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _read_pdf_pages(path: Path) -> list[tuple[int, str]]:
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(str(path))
        pages = []
        try:
            for i, page in enumerate(pdf, start=1):
                tp = page.get_textpage()
                pages.append((i, tp.get_text_bounded()))
        finally:
            pdf.close()
        return pages
    except Exception:
        # Fallback: pypdf (pure python, slightly different extraction quality)
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [(i + 1, (page.extract_text() or "")) for i, page in enumerate(reader.pages)]
