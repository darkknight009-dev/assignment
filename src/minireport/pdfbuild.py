"""Assemble the combined PDF: report typeset + Excel table + PDF attachments.

Word->PDF conversion is out of scope for a lightweight prototype, so the
report is re-typeset with ReportLab from the exact same SectionContent and
Evidence objects the DOCX was written from; DOCX and combined PDF therefore
cannot disagree.  The Excel results table is rendered from the generated
workbook itself (raw cells read back via openpyxl), giving a readable
rendering of the Excel table, and the original supporting PDF attachments are
appended verbatim after labelled separator pages.
"""

from __future__ import annotations

from pathlib import Path

from .models import DRAFT_BANNER, Evidence, Protocol, SectionContent


def read_xlsx_rows(xlsx_path: Path) -> list[list[str]]:
    """Read back the generated workbook so the PDF renders the real cells."""
    from openpyxl import load_workbook

    wb = load_workbook(str(xlsx_path), data_only=True)
    rows: list[list[str]] = []
    for ws in wb.worksheets:
        rows.append([f"[Sheet: {ws.title}]"])
        for row in ws.iter_rows():
            values = [("" if c.value is None else str(c.value)) for c in row]
            if any(v.strip() for v in values):
                rows.append(values)
        rows.append([""])
    return rows


def build_combined_pdf(sections: list[SectionContent], evidence: Evidence,
                       protocol: Protocol, xlsx_rows: list[list[str]],
                       attachments_by_label: dict[str, str],
                       out_path: Path) -> None:
    """Typeset the report, render the Excel table, then append attachments."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=13.5,
                        spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#1a3d5c"))
    process = ParagraphStyle("Process", parent=styles["BodyText"], fontSize=8.5,
                             fontName="Helvetica-Oblique", textColor=colors.HexColor("#555555"),
                             spaceAfter=4)
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=9.5, leading=12.5, spaceAfter=3)
    bullet = ParagraphStyle("Bullet", parent=body, leftIndent=14)
    note = ParagraphStyle("Note", parent=body, fontName="Helvetica-Oblique",
                          textColor=colors.HexColor("#996600"), fontSize=8.5)
    banner = ParagraphStyle("Banner", parent=styles["Title"], fontSize=14,
                            textColor=colors.HexColor("#9C0006"), alignment=TA_CENTER)

    def esc(text: str) -> str:
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story: list = []

    # ------------------------------------------------------------- cover -----
    story.append(Paragraph(DRAFT_BANNER, banner))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(f"Protocol {esc(protocol.protocol_id)} - {esc(protocol.component)}",
                           ParagraphStyle("Sub", parent=styles["Title"], fontSize=13)))
    story.append(Spacer(1, 3 * mm))

    # ----------------------------------------------------------- sections ----
    for sec in sections:
        story.append(Paragraph(esc(sec.key), h1))
        story.append(Paragraph(f"[Process control: {esc(process_label(sec))}]", process))
        for raw in sec.body.split("\n"):
            line = raw.rstrip()
            if not line.strip():
                continue
            if line.startswith("- "):
                story.append(Paragraph("&bull; " + esc(line[2:]), bullet))
            else:
                story.append(Paragraph(esc(line), body))
        for w in sec.warnings:
            story.append(Paragraph("Note: " + esc(w), note))

        # results section additionally renders the results table (same numbers
        # as the DOCX table and the XLSX: all from one Evidence object)
        if sec.key == "Results and evidence":
            story.append(Spacer(1, 2 * mm))
            story.extend(_results_table(evidence))

        if sec.citations:
            story.append(Paragraph("Sources: " + "; ".join(esc(c.reference) for c in sec.citations), note))
        story.append(Spacer(1, 3 * mm))

    # ---------------------------------------------------- excel rendering ----
    story.append(PageBreak())
    story.append(Paragraph("Generated Excel table (results_table.xlsx)", h1))
    story.append(Paragraph("Readable rendering of the generated workbook - identical "
                           "values to the attached XLSX file.", body))
    story.extend(_render_generic_table(xlsx_rows))

    doc = SimpleDocTemplate(str(out_path), pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title=f"Draft - Requires Review - Protocol {protocol.protocol_id}")
    doc.build(story)

    # ------------------------------------------------- append attachments ----
    if attachments_by_label:
        _append_pdfs(attachments_by_label, out_path)


def process_label(sec: SectionContent) -> str:
    label = {
        "deterministic": "deterministic / template-based",
        "rag": "RAG - retrieval over reference documents, sources cited",
        "evidence": "evidence-based, deterministic processing",
        "ai": "AI-assisted narrative (fact-checked)",
        "deterministic (AI unavailable)": "deterministic (AI unavailable; see notes)",
        "deterministic (AI draft discarded)": "deterministic (AI draft discarded on fact check; see notes)",
    }
    return label.get(sec.process) or label.get(sec.process.split(" ")[0], sec.process)


def _results_table(evidence: Evidence):
    from reportlab.lib import colors
    from reportlab.platypus import Spacer, Table, TableStyle

    data = [["Sample ID", "Measurement", "Unit", "Verdict", "CSV row"]]
    for ev in evidence.evaluations:
        data.append([ev["sample_id"], f"{ev['measurement_display']} {ev['unit']}".strip(),
                     ev["unit"], ev["verdict"], f"row {ev['row_number']}"])
    t = Table(data, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3d5c")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7fa")]),
    ]))
    _highlight_fails(t, evidence, TableStyle)
    return [Spacer(1, 2), t, Spacer(1, 4)]


def _highlight_fails(table, evidence: Evidence, TableStyle) -> None:
    from reportlab.lib import colors

    for i, ev in enumerate(evidence.evaluations, start=1):
        style_cmds = []
        if ev["verdict"] == "Fail":
            style_cmds = [
                ("TEXTCOLOR", (0, i), (-1, i), colors.HexColor("#9C0006")),
                ("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"),
                ("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FDECEA")),
            ]
        elif ev["verdict"] == "NO DATA":
            style_cmds = [("TEXTCOLOR", (0, i), (-1, i), colors.HexColor("#996600"))]
        if style_cmds:
            table.setStyle(TableStyle(style_cmds))


def _render_generic_table(rows: list[list[str]]):
    from reportlab.lib import colors
    from reportlab.platypus import Spacer, Table, TableStyle

    # filter sheet-marker rows to keep tables rectangular
    table_rows = [r for r in rows if not (len(r) == 1 and r[0].startswith("[Sheet:")) and len(r) >= 1]
    if not table_rows or not any(any(c.strip() for c in r) for r in table_rows):
        return []
    width = max(len(r) for r in table_rows)
    normalised = [r + [""] * (width - len(r)) for r in table_rows]
    t = Table(normalised, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3d5c")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7fa")]),
    ]))
    return [Spacer(1, 2), t]


def _append_pdfs(attachments_by_label: dict[str, str], out_path: Path) -> None:
    """Append each supplied PDF after a labelled separator page.

    Separator pages carry visible "Attachment A - filename" text so a
    printout can be navigated without a PDF viewer.
    """
    import io

    from pypdf import PdfReader, PdfWriter
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    cover = ParagraphStyle("Cover", parent=styles["Title"], alignment=TA_CENTER)
    small = ParagraphStyle("Small", parent=styles["BodyText"], alignment=TA_CENTER,
                           textColor=colors.HexColor("#555555"))

    writer = PdfWriter()
    main_reader = PdfReader(str(out_path))
    for page in main_reader.pages:
        writer.add_page(page)

    for label in sorted(attachments_by_label):
        path = Path(attachments_by_label[label])
        try:
            reader = PdfReader(str(path))
            pages = list(reader.pages)
        except Exception:
            pages = []  # unreadable attachment: keep the separator, skip pages

        buf = io.BytesIO()
        SimpleDocTemplate(buf, pagesize=A4).build([
            Paragraph(f"Attachment {label.split()[-1]}", cover),
            Spacer(1, 2 * mm),
            Paragraph(path.name, small),
        ])
        for page in PdfReader(io.BytesIO(buf.getvalue())).pages:
            writer.add_page(page)
        for page in pages:
            writer.add_page(page)

    with open(out_path, "wb") as fh:
        writer.write(fh)
